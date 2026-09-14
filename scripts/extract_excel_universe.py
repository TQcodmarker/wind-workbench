"""Read the supplied workbook's query universe without changing the workbook.

Only bond identifiers and their cell references are extracted. Excel's cached
financial values must not be treated as observations for the requested date.
"""
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re

import openpyxl


CODE = re.compile(r"^\d{6,9}\.(?:IB|SH|SZ)$")


def digest(path: Path) -> str:
    with path.open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def extract(source: Path, copy: Path | None, target: str) -> dict:
    source_hash = digest(source)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["原数据-WIND"]
    references: dict[str, list[str]] = defaultdict(list)
    invalid = []
    row_count = 0
    for row in sheet.iter_rows(min_row=2, max_col=1):
        cell = row[0]
        row_count += 1
        value = "" if cell.value is None else str(cell.value).strip()
        code = value.upper()
        if not CODE.fullmatch(code):
            invalid.append({"row": cell.row, "cell": cell.coordinate, "value": value})
            continue
        references[code].append(cell.coordinate)
    duplicate_rows = [
        {"code": code, "cells": cells, "rows": [int(cell[1:]) for cell in cells]}
        for code, cells in references.items() if len(cells) > 1
    ]
    per_sheet = {}
    all_sheet_codes = set()
    for candidate in workbook:
        found = set()
        occurrences = 0
        for row in candidate.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str) and CODE.fullmatch(value.strip().upper()):
                    found.add(value.strip().upper())
                    occurrences += 1
        all_sheet_codes.update(found)
        per_sheet[candidate.title] = {"uniqueCodes": len(found), "occurrences": occurrences}
    header = sheet["A1"].value
    if isinstance(header, (date, datetime)):
        header = header.isoformat()
    result = {
        "schemaVersion": 1,
        "targetDate": target,
        "sourcePath": str(source.resolve()),
        "sourceHash": source_hash,
        "sourceHashAlgorithm": "sha256",
        "workspaceCopyPath": str(copy.resolve()) if copy and copy.exists() else None,
        "workspaceCopyHash": digest(copy) if copy and copy.exists() else None,
        "sheet": sheet.title,
        "codeColumn": "A",
        "codeColumnHeader": header,
        "range": f"A2:A{sheet.max_row}",
        "rowCount": row_count,
        "uniqueCodes": len(references),
        "invalidRows": invalid,
        "duplicateRows": duplicate_rows,
        "codes": list(references),
        "rowRefs": dict(references),
        "allWorksheetCodeCounts": per_sheet,
        "codesInOtherSheetsNotInSelectedSheet": sorted(all_sheet_codes - references.keys()),
        "queryScope": "all_valid_unique_codes_in_selected_sheet",
        "filtersApplied": [],
        "cachedValuesUsedAsMarketData": False,
    }
    result["workspaceCopyIdentical"] = result["workspaceCopyHash"] == source_hash
    workbook.close()
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--copy", type=Path)
    parser.add_argument("--date", required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    date.fromisoformat(args.date)
    result = extract(args.source, args.copy, args.date)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({key: value for key, value in result.items() if key not in {"codes", "rowRefs"}}, ensure_ascii=False))


if __name__ == "__main__":
    main()
